from docx import Document
from langchain_core.documents import Document as LCDocument

def load_docx(file):
    doc = Document(file)
    text = "\n".join([p.text for p in doc.paragraphs])

    return LCDocument(page_content=text)